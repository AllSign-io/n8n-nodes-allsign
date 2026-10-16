#!/usr/bin/env python3
"""Genera el acuerdo de confidencialidad en español, en DOCX, con las variables
que la plantilla de HubSpot puede llenar sola.

Uso:  python3 examples/generate_nda_es.py [ruta-de-salida.docx]
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

AZUL = RGBColor(0x00, 0x7B, 0xFF)
TINTA = RGBColor(0x1A, 0x1A, 0x2E)
GRIS = RGBColor(0x55, 0x55, 0x55)

# Junto al script, no junto a donde lo corriste: si no, correrlo desde
# examples/ truena con un error de archivo que no dice nada.
SALIDA = sys.argv[1] if len(sys.argv) > 1 else str(
    Path(__file__).resolve().parent / "NDA_Plantilla_AllSign_ES.docx")


def sombrear(celda, color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = etree.SubElement(tc_pr, qn("w:shd"))
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")


def construir():
    doc = Document()

    for seccion in doc.sections:
        seccion.top_margin = Cm(2.5)
        seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = Cm(2.5)
        seccion.right_margin = Cm(2.5)

    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(11)
    normal.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Sin esto Word corrige el documento en inglés y subraya en rojo todo el texto.
    idioma = doc.styles["Normal"].element.get_or_add_rPr()
    lang = idioma.makeelement(qn("w:lang"), {qn("w:val"): "es-MX"})
    idioma.append(lang)

    props = doc.core_properties
    props.title = "Acuerdo de confidencialidad"
    props.author = "AllSign"
    props.language = "es-MX"

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(4)
    run = titulo.add_run("ACUERDO DE CONFIDENCIALIDAD")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = TINTA

    bajada = doc.add_paragraph()
    bajada.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bajada.paragraph_format.space_after = Pt(14)
    run = bajada.add_run("Entre las partes que se identifican a continuación")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = GRIS

    datos = doc.add_table(rows=4, cols=2)
    datos.style = "Table Grid"
    filas = [
        ("Fecha de celebración", "{{ fecha_efectiva }}"),
        ("Empresa", "{{ empresa_emisora }}"),
        ("Representante del cliente", "{{ nombre_cliente }}"),
        ("Empresa del cliente", "{{ empresa_cliente }}"),
    ]
    for i, (etiqueta, valor) in enumerate(filas):
        celda = datos.cell(i, 0)
        sombrear(celda, "F4F6F8")
        run = celda.paragraphs[0].add_run(etiqueta)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = GRIS

        run = datos.cell(i, 1).paragraphs[0].add_run(valor)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = AZUL

    doc.add_paragraph()

    def encabezado(numero, texto, pegado=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = pegado
        run = p.add_run(f"{numero}. {texto}")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = TINTA

    def cuerpo(texto, pegado=False):
        p = doc.add_paragraph(texto)
        p.paragraph_format.line_spacing = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        # Una hoja de firmas suelta se puede separar y pegar a otro documento. Desde la
        # cláusula 10 todo viaja junto para que caiga en la misma página.
        p.paragraph_format.keep_with_next = pegado
        return p

    def vinetas(items):
        for item in items:
            p = doc.add_paragraph(item, style="List Bullet")
            p.paragraph_format.line_spacing = Pt(15)

    encabezado(1, "OBJETO Y PARTES")
    cuerpo(
        "Este acuerdo se celebra el {{ fecha_efectiva }} entre {{ empresa_emisora }} y "
        "{{ empresa_cliente }}, representada en este acto por {{ nombre_cliente }}. Ambas "
        "empresas se denominan conjuntamente las Partes, y cada una queda obligada en los "
        "términos de este acuerdo."
    )
    cuerpo(
        "Quien firma por cada Parte declara contar con las facultades necesarias para "
        "representarla y obligarla, y que dichas facultades no le han sido revocadas ni "
        "limitadas a la fecha de firma."
    )
    cuerpo(
        "El objeto de este acuerdo es proteger la información confidencial que las Partes se "
        "compartan con motivo de {{ objeto }}. El acuerdo es recíproco: cada Parte actúa como "
        "Parte Reveladora respecto de la información que comparte, y como Parte Receptora "
        "respecto de la que recibe de la otra."
    )

    encabezado(2, "QUÉ SE CONSIDERA INFORMACIÓN CONFIDENCIAL")
    cuerpo(
        "Toda información, en cualquier forma o medio, que una Parte comparta con la otra "
        "y que esté marcada como confidencial, o que por su naturaleza y por las circunstancias "
        "en que se comparte deba entenderse como tal. Queda comprendida, de manera "
        "enunciativa y no limitativa:"
    )
    vinetas([
        "Planes de negocio, estrategias e información financiera",
        "Datos técnicos, secretos industriales y conocimiento especializado",
        "Planes de producto, diseños y especificaciones",
        "Listas de clientes, estrategias de mercadotecnia e información de ventas",
        "Programas de cómputo, código fuente, algoritmos y bases de datos",
        "Datos personales a los que se tenga acceso con motivo de este acuerdo",
    ])

    encabezado(3, "OBLIGACIONES DE CADA PARTE")
    cuerpo(
        "Cada Parte, respecto de la información confidencial que reciba de la otra, se obliga a:"
    )
    vinetas([
        "Guardarla con estricta reserva",
        "No revelarla a terceros sin autorización previa y por escrito de la Parte que la "
        "compartió, salvo a las personas a que se refiere la última viñeta de esta cláusula",
        "Usarla únicamente para el fin descrito en la cláusula primera",
        "Protegerla con el mismo cuidado con que protege la propia, y nunca con menos "
        "del que resulta razonable",
        "Avisar a la otra Parte en cuanto tenga conocimiento de cualquier divulgación o "
        "uso no autorizado",
        "Limitar el acceso a las personas que necesiten conocerla y que estén sujetas a "
        "obligaciones de confidencialidad equivalentes",
    ])

    encabezado(4, "TRATAMIENTO DE DATOS PERSONALES")
    cuerpo(
        "Cuando la información confidencial incluya datos personales, la Parte que los "
        "reciba los tratará conforme a la Ley Federal de Protección de Datos Personales en "
        "Posesión de los Particulares vigente y su normativa aplicable, los usará solo para "
        "el fin de este acuerdo, y los suprimirá o devolverá cuando deje de ser necesario "
        "conservarlos. Esta obligación aplica a cada Parte respecto de los datos personales "
        "que reciba de la otra."
    )

    encabezado(5, "EXCEPCIONES")
    cuerpo("Este acuerdo no aplica a la información que:")
    vinetas([
        "Ya era del dominio público al momento de compartirse, o llegó a serlo después sin "
        "culpa de la Parte que la recibió",
        "La Parte que la recibió ya conocía legítimamente y sin obligación de reserva",
        "Un tercero compartió con la Parte que la recibió sin violar ninguna obligación de "
        "confidencialidad",
        "La Parte que la recibió desarrolló por su cuenta sin usar la información confidencial",
    ])
    cuerpo(
        "La Parte que deba revelar información confidencial por mandato de autoridad competente "
        "lo hará solo en la medida requerida, avisando a la otra Parte con la anticipación que "
        "le sea posible. Para cualquier otro fin, esa información sigue siendo confidencial."
    )

    encabezado(6, "VIGENCIA DEL ACUERDO Y DE LA CONFIDENCIALIDAD")
    cuerpo(
        "Este acuerdo entra en vigor el {{ fecha_efectiva }} y permanece vigente hasta que "
        "cualquiera de las Partes lo dé por terminado, avisando a la otra por escrito con "
        "treinta días naturales de anticipación."
    )
    cuerpo(
        "Las obligaciones de confidencialidad subsisten por un plazo de {{ vigencia }}, "
        "contado a partir de la terminación de este acuerdo, y alcanzan a toda la información "
        "compartida durante su vigencia. Tratándose de información que constituya secreto "
        "industrial, subsisten mientras conserve ese carácter conforme a la ley, y tratándose "
        "de datos personales, subsisten conforme a la ley de la materia."
    )

    encabezado(7, "DEVOLUCIÓN DE LA INFORMACIÓN")
    cuerpo(
        "Terminado el acuerdo conforme a la cláusula anterior, o cuando la Parte que "
        "compartió la información lo solicite, la Parte que la recibió la devolverá o "
        "destruirá, incluidas las copias, y lo confirmará por escrito. Puede conservar una "
        "copia únicamente cuando una obligación legal se lo exija, y seguirá protegiéndola "
        "en los términos de este acuerdo mientras la conserve."
    )

    encabezado(8, "INCUMPLIMIENTO")
    cuerpo(
        "El incumplimiento de estas obligaciones da derecho a la Parte afectada a exigir el "
        "pago de los daños y perjuicios que se le hayan causado, sin perjuicio de las demás "
        "acciones que le correspondan."
    )

    encabezado(9, "LEY APLICABLE Y TRIBUNALES COMPETENTES")
    cuerpo(
        "Este acuerdo es de naturaleza mercantil y se rige por las leyes federales de los "
        "Estados Unidos Mexicanos."
    )
    cuerpo(
        "Para su interpretación y cumplimiento, las Partes se someten a los tribunales "
        "competentes de {{ jurisdiccion }} y renuncian a cualquier otro fuero que pudiera "
        "corresponderles por razón de sus domicilios presentes o futuros."
    )

    encabezado(10, "FIRMA ELECTRÓNICA", pegado=True)
    cuerpo(
        "Las Partes aceptan celebrar y firmar este acuerdo por medios electrónicos, en los "
        "términos del Código de Comercio, y reconocen que el mensaje de datos que lo "
        "contiene les es atribuible.", pegado=True
    )
    cuerpo(
        "La constancia de conservación emitida conforme a la NOM-151 acredita la "
        "conservación e integridad del documento desde la fecha y hora asentadas en la propia "
        "constancia, y no la validez de la "
        "firma, que se rige por las disposiciones aplicables.", pegado=True
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    run = p.add_run("FIRMAS")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = TINTA

    cuerpo(
        "Leído el contenido de este acuerdo y enteradas las Partes de su alcance y fuerza "
        "legal, lo firman de conformidad quienes las representan.", pegado=True
    )

    firmas = doc.add_table(rows=2, cols=2)
    firmas.autofit = True
    # Las etiquetas de abajo son el ancla de las cajas de firma en AllSign, así que cada
    # una tiene que aparecer UNA sola vez en todo el documento. Por eso no se repiten las
    # de la tabla de partes.
    for columna, (quien, quien_detalle) in enumerate([
        ("{{ empresa_emisora }}",
         "Por conducto de {{ representante_emisor }}\nFirma por la empresa"),
        ("{{ empresa_cliente }}",
         "Por conducto de {{ nombre_cliente }}  ·  {{ correo_cliente }}\nFirma por el cliente"),
    ]):
        celda = firmas.cell(0, columna)
        celda.paragraphs[0].add_run("\n\n\n_______________________________")
        p = firmas.cell(1, columna).paragraphs[0]
        run = p.add_run(quien)
        run.bold = True
        run.font.size = Pt(10)
        p.add_run("\n")
        run = p.add_run(quien_detalle)
        run.font.size = Pt(9)
        run.font.color.rgb = GRIS

    doc.save(SALIDA)
    return SALIDA


if __name__ == "__main__":
    print(construir())
